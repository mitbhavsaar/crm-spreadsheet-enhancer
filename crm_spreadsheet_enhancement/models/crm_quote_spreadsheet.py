# -*- coding: utf-8 -*-
from odoo import api, fields, models, _
import json
import logging
import re
import openpyxl
import uuid
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import range_boundaries
from openpyxl.formula.translate import Translator

_logger = logging.getLogger(__name__)

CRM_MATERIAL_LINE_BASE_FIELDS = [
    'product_template_id',
    'quantity',
    'price',
]

class CrmLeadSpreadsheet(models.Model):
    _name = 'crm.lead.spreadsheet'
    _inherit = 'spreadsheet.mixin'
    _description = 'CRM Quotation Spreadsheet'

    name = fields.Char(required=True)
    lead_id = fields.Many2one('crm.lead', string="Opportunity", ondelete='cascade')
    sale_id = fields.Many2one('sale.order', string="Sale Order", ondelete='set null')
    company_id = fields.Many2one('res.company', default=lambda self: self.env.company)
    raw_spreadsheet_data = fields.Text("Raw Spreadsheet Data")

    def _get_line_root_info(self, line):
        root = line.original_line_id or line
        return f"sheet_{root.id}", str(root.id)

    # ------------------------------------------------------------------
    # ✅ CRITICAL: Override get_list_data (PUBLIC METHOD)
    # ------------------------------------------------------------------
    @api.model
    def get_list_data(self, model, list_id, field_names):
        """
        ✅ THIS IS CALLED BY SPREADSHEET JS
        Override base spreadsheet method to handle dynamic attributes
        """
        
        if model != 'crm.material.line':
            return super().get_list_data(model, list_id, field_names)
        
        try:
            root_id = int(list_id)
        except (ValueError, TypeError):
            _logger.error(f"❌ Invalid list_id: {list_id}")
            return []

        # ✅ PARENT-CHILD Logic:
        # If root_id corresponds to a Parent line, find all its Child lines.
        # If it's just a regular line, use it as is.
        lines = self.env['crm.material.line'].search([
            ('original_line_id', '=', root_id),
            ('is_dimension_row', '=', True)
        ], order='id asc')

        if not lines:
            # Fallback to the root itself if no children found (for backward compatibility)
            lines = self.env['crm.material.line'].search([
                ('id', '=', root_id),
                ('is_dimension_row', '=', False)
            ])

        if not lines:
            _logger.warning(f"❌ No material lines found for list {list_id}")
            return []

        results = []
        for line in lines:
            # Get attributes_json
            attrs = line.attributes_json or {}
            row = {"id": line.id}
            
            # 🔍 DEBUG: Log raw attributes_json for Length/Width
            if 'Length' in attrs or 'Width' in attrs:
                _logger.info(f"📊 [RAW JSON] Line {line.id}: attributes_json = {attrs}")

            for field in field_names:
                # 🔥 Grouping logic for multiple lengths/dimension rows
                # If this is a secondary dimension row, hide header fields for a cleaner grouped look
                grouping_fields = [
                    'product_template_id', 'product_display_name', 
                    'quantity', 'product_uom_id', 'Quantity UOM',
                    'price', 'quoted_price', 'total_price'
                ]
                if line.is_dimension_row and field in grouping_fields:
                    row[field] = ""
                    continue

                if field in line._fields:
                    # Standard Odoo field
                    val = line[field]
                    if hasattr(val, "display_name"):
                        row[field] = val.display_name
                    else:
                        row[field] = val
                    
                    # 🔍 DEBUG: Log Length/Width from model fields
                    if field in ['length', 'width']:
                        _logger.info(f"📊 [Sheet Data] Line {line.id}: {field} = {val} (from MODEL FIELD)")
                else:
                    # Dynamic attribute from attributes_json
                    val = attrs.get(field, "")
                    
                    # � DEBUG: Log Length/Width from attributes_json
                    if field in ['Length', 'Width']:
                        _logger.info(f"📊 [Sheet Data] Line {line.id}: {field} = '{val}' (from ATTRIBUTES_JSON)")
                    
                    # �🔥 Strip unit suffix from Length and Width columns
                    # Show only numeric value (e.g., "100" instead of "100 mm")
                    if field in ['Length', 'Width'] and val:
                        try:
                            import re
                            # Extract only the numeric part
                            numeric_match = re.search(r'[-+]?\d*\.?\d+', str(val))
                            if numeric_match:
                                numeric_value = float(numeric_match.group())
                                row[field] = numeric_value
                                _logger.info(f"   ✅ Stripped to numeric: {numeric_value}")
                            else:
                                row[field] = val
                                _logger.info(f"   ⚠️ No numeric match, using original: {val}")
                        except Exception as e:
                            row[field] = val
                            _logger.warning(f"   ❌ Error extracting numeric: {e}")
                    else:
                        row[field] = val
            
            results.append(row)
        
        return results

    def write(self, vals):
        """✅ Defeat Reversion: Push to records on EVERY save"""
        res = super().write(vals)
        if 'raw_spreadsheet_data' in vals:
            _logger.info(f"💾 [DEBUGGER] CRM SAVE TRIGGERED for Spreadsheet {self.id}. Pushing to records...")
            for record in self:
                try:
                    record._push_spreadsheet_data_to_records()
                    _logger.info(f"✅ [DEBUGGER] Data successfully pushed to Material Lines for Spreadsheet {record.id}")
                except Exception as e:
                    _logger.error(f"❌ [DEBUGGER] CRM Force push failed: {e}")
        return res

    def _push_spreadsheet_data_to_records(self):
        """✅ Manual PUSH with FORMULA SHIELD for CRM"""
        self.ensure_one()
        if not self.raw_spreadsheet_data: return
        
        try:
            data = json.loads(self.raw_spreadsheet_data)
            lists = data.get('lists', {})
            sheets = data.get('sheets', [])
            
            # PHASE 1: Build a map of listId -> sorted record IDs
            # This allows us to resolve which record is at which row.
            list_records_map = {}
            for list_id, list_meta in lists.items():
                domain = list_meta.get('domain', [])
                root_id = None
                # Check for id = X
                for cond in domain:
                    if isinstance(cond, list) and len(cond) >= 3 and cond[0] == 'id' and cond[1] == '=':
                        root_id = cond[2]
                        break
                
                if not root_id and str(list_id).isdigit():
                    root_id = int(list_id)

                if root_id:
                    # ✅ PARENT-CHILD Logic:
                    # Find all Child lines for this Parent
                    group_lines = self.env['crm.material.line'].sudo().search([
                        ('original_line_id', '=', root_id),
                        ('is_dimension_row', '=', True)
                    ], order='id asc')
                    
                    if not group_lines:
                         # Fallback
                         group_lines = self.env['crm.material.line'].sudo().search([
                            ('id', '=', root_id)
                        ])

                    list_records_map[str(list_id)] = group_lines.ids
                    _logger.debug(f"📊 [CRM Push] List {list_id} mapped to records: {group_lines.ids}")

            updates = {}

            # Helper to get row number from cell ref (A1 -> 0, B2 -> 1)
            def get_row_index(ref):
                match = re.search(r'\d+', ref)
                return int(match.group()) - 1 if match else -1

            for sheet in sheets:
                field_syncs = sheet.get('fieldSyncs', {})
                cells = sheet.get('cells', {})
                
                for cell_ref, sync in field_syncs.items():
                    list_id_str = str(sync.get('listId', ''))
                    field_name = sync.get('fieldName')
                    
                    if not list_id_str or not field_name: continue

                    # Resolve record ID based on row position in the Odoo list
                    row_idx = get_row_index(cell_ref)
                    if row_idx < 1: continue # Skip header or invalid

                    # record_index = row_idx - table_start_row
                    # Since we insert lists at row 0 (header), record 0 is row 1.
                    record_idx = row_idx - 1
                    
                    record_ids = list_records_map.get(list_id_str, [])
                    if record_idx < 0 or record_idx >= len(record_ids):
                        _logger.warning(f"⚠️ [CRM Push] Cell {cell_ref} row {row_idx} is out of bounds for list {list_id_str} (size {len(record_ids)})")
                        continue
                    
                    record_id = record_ids[record_idx]
                    
                    # ✅ Try to get the evaluated value
                    raw_value = sync.get('value')
                    if raw_value is None:
                        cell = cells.get(cell_ref, {})
                        raw_value = cell.get('content', '')
                        if str(raw_value).startswith('='): continue

                    if record_id not in updates:
                        updates[record_id] = {}
                    updates[record_id][field_name] = raw_value
                    _logger.info(f"📍 [CRM Push] Queueing update: Line {record_id}, Field {field_name}, Value {raw_value}")

            # Apply updates
            for record_id, vals in updates.items():
                line = self.env['crm.material.line'].browse(record_id)
                if not line.exists() or line.lead_id != self.lead_id: continue
                
                processed_vals = {}
                for f, v in vals.items():
                    if f in line._fields:
                        field = line._fields[f]
                        if field.compute and not field.inverse: continue
                        
                        field_type = field.type
                        try:
                            # Clean up value: remove currency symbols, spaces, and handle grouping separators
                            clean_v = str(v).replace('$', '').replace('€', '').replace('£', '').replace('₹', '').replace(',', '').strip() if v not in (None, False, '') else '0'
                            
                            if field_type in ['float', 'monetary']:
                                processed_vals[f] = float(clean_v) if clean_v else 0.0
                            elif field_type == 'integer':
                                processed_vals[f] = int(float(clean_v)) if clean_v else 0
                            elif field_type == 'boolean':
                                processed_vals[f] = str(v).lower() in ('true', '1', 'yes')
                            elif field_type == 'many2one':
                                if v and not str(v).isdigit():
                                    comodel = field.comodel_name
                                    target = self.env[comodel].search([('display_name', '=', v)], limit=1)
                                    if target: processed_vals[f] = target.id
                                else:
                                    processed_vals[f] = int(v) if v and str(v).isdigit() else False
                            else:
                                processed_vals[f] = v
                        except: continue
                    else:
                        processed_vals[f] = v
                
                if processed_vals:
                    line.sudo().write(processed_vals)

        except Exception as e:
            _logger.error(f"❌ CRM PUSH Sync failed: {e}")

    def _update_spreadsheet_lists_data(self, check_only=False):
        """
        ✅ SMART-SYNC: Refresh spreadsheet cells if core fields changed.
        If check_only=True, only returns True if a mismatch is found in CORE fields.
        """
        self.ensure_one()
        if not self.raw_spreadsheet_data: return False
        
        try:
            data = json.loads(self.raw_spreadsheet_data)
            lists = data.get('lists', {})
            sheets = data.get('sheets', [])
            modified = False

            for sheet in sheets:
                field_syncs = sheet.get('fieldSyncs', {})
                cells = sheet.get('cells', {})
                
                for cell_ref, sync in field_syncs.items():
                    list_id = sync.get('listId', '')
                    field_name = sync.get('fieldName')
                    
                    record_id = None
                    if list_id in lists:
                        domain = lists[list_id].get('domain', [])
                        for cond in domain:
                            if isinstance(cond, list) and len(cond) >= 3 and cond[0] == 'id' and cond[1] == '=':
                                record_id = cond[2]
                                break
                    if not record_id and str(list_id).isdigit():
                        record_id = int(list_id)

                    if record_id and field_name:
                        line = self.env['crm.material.line'].browse(record_id)
                        if line.exists():
                            val = None
                            if field_name in line._fields:
                                val = line[field_name]
                            else:
                                attrs = line.attributes_json or {}
                                val = attrs.get(field_name)

                            # 🔥 FIX: Strip units from Length/Width for direct cell sync
                            if field_name in ['Length', 'Width'] and val:
                                try:
                                    import re
                                    numeric_match = re.search(r'[-+]?\d*\.?\d+', str(val))
                                    if numeric_match:
                                        val = float(numeric_match.group())
                                        _logger.info(f"   ✅ [Sync] Stripped {field_name} unit: {val}")
                                except Exception:
                                    pass

                            if val is not None:
                                new_val = val.display_name if hasattr(val, "display_name") else val
                                if new_val is False: new_val = ''
                                
                                if cell_ref not in cells: cells[cell_ref] = {}
                                current_val = cells[cell_ref].get('content', '')

                                # ✅ FORMULA SHIELD
                                if str(current_val).startswith('='): continue

                                # ✅ SMART-SYNC CHECK
                                if check_only:
                                    # Ignore resin fields for core comparison
                                    is_resin_field = any(x in str(field_name).lower() for x in ['resin', 'rsign', 'raisin', 'resign'])
                                    if not is_resin_field and str(current_val) != str(new_val):
                                        _logger.info(f"🚨 [CRM SMART-SYNC] Core mismatch: {field_name}")
                                        return True
                                    continue

                                if str(current_val) != str(new_val):
                                    cells[cell_ref]['content'] = str(new_val)
                                    modified = True

            if check_only: return False

            if modified:
                self.raw_spreadsheet_data = json.dumps(data)
                return True
                
        except Exception as e:
            _logger.error(f"❌ CRM PULL Sync failed: {e}")
            return False
        
        return False

    # ------------------------------------------------------------------
    # ✅ INTERNAL: _get_list_data (PRIVATE METHOD)
    # ------------------------------------------------------------------
    def _get_list_data(self, list_id):
        """
        Internal method for other operations
        """
        self.ensure_one()
        
        try:
            list_id_int = int(list_id)
        except (ValueError, TypeError):
            return []

        line = self.env['crm.material.line'].browse(list_id_int)
        if not line.exists():
            return []

        row = {
            'id': line.id,
            'product_template_id': line.product_template_id.display_name if line.product_template_id else '',
            'quantity': line.quantity or 0,
        }

        attrs = line.attributes_json or {}
        for key, value in attrs.items():
            row[key] = value

        return [row]

    # ------------------------------------------------------------------
    # OPEN FORMVIEW
    # ------------------------------------------------------------------
    def get_formview_action(self, access_uid=None):
        return self.action_open_spreadsheet()

    def action_open_spreadsheet(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.client',
            'tag': 'action_crm_lead_spreadsheet',
            'params': {
                'spreadsheet_id': self.id,
                'model': 'crm.lead.spreadsheet',
                'lead_id': self.lead_id.id if self.lead_id else False,
                'lead_display_name': self.lead_id.name if self.lead_id else False,
            },
        }

    # ------------------------------------------------------------------
    # CREATE
    # ------------------------------------------------------------------
    @api.model_create_multi
    def create(self, vals_list):
        records = super().create(vals_list)
        for rec in records:
            if rec.lead_id and rec.lead_id.material_line_ids:
                for line in rec.lead_id.material_line_ids:
                    rec.with_context(material_line_id=line.id)._dispatch_insert_list_revision()
        return records

    # ------------------------------------------------------------------
    # JOIN SESSION
    # ------------------------------------------------------------------
    def join_spreadsheet_session(self, access_token=None):
        self.ensure_one()

        # ✅ SMART-SYNC: Only refresh if core fields changed
        _logger.info("🛠️ [CRM SMART-SYNC] Checking if core fields have changed...")
        try:
            if self._update_spreadsheet_lists_data(check_only=True):
                _logger.info("🚨 [CRM SMART-SYNC] Core mismatch. Refreshing...")
                self._update_spreadsheet_lists_data(check_only=False)
            else:
                _logger.info("✅ [CRM SMART-SYNC] All core fields match. Skipping refresh.")
        except Exception as e:
            _logger.error(f"❌ CRM Smart Sync error: {e}")
        
        self._sync_sheets_with_material_lines()

        data = super().join_spreadsheet_session(access_token)
        data.update({
            'lead_id': self.lead_id.id if self.lead_id else False,
            'lead_display_name': self.lead_id.display_name if self.lead_id else False,
            'spreadsheet_id': self.id,
            'sheet_id': self.id
        })

        spreadsheet_json = data.get('data') or {}
        lists = spreadsheet_json.get('lists') or {}
        sheets = spreadsheet_json.get('sheets') or []

        # ✅ Group current lines by root for correct sheet mapping
        root_to_lines = {}
        for line in self.lead_id.material_line_ids:
            _, root_id_str = self._get_line_root_info(line)
            root_id = int(root_id_str)
            if root_id not in root_to_lines:
                root_to_lines[root_id] = []
            root_to_lines[root_id].append(line.id)

        # Update existing lists with correct domains (including duplicates/children)
        for list_id, list_config in lists.items():
            if list_id.isdigit():
                root_id = int(list_id)
                # Find children for this root/parent
                children_ids = self.env['crm.material.line'].sudo().search([
                    ('original_line_id', '=', root_id),
                    ('is_dimension_row', '=', True)
                ], order='id asc').ids
                
                if children_ids:
                    list_config['domain'] = [['id', 'in', children_ids]]
                    list_config['orderBy'] = [{'name': 'id', 'desc': False}]
                elif root_id in root_to_lines:
                     # Fallback for legacy
                     list_config['domain'] = [['id', 'in', root_to_lines[root_id]]]
                     list_config['orderBy'] = [{'name': 'id', 'desc': False}]

        existing_list_ids = {int(list_id) for list_id in lists.keys() if list_id.isdigit()}
        missing_root_ids = set(root_to_lines.keys()) - existing_list_ids
        removed_ids = existing_list_ids - set(root_to_lines.keys())

        # Add sheets for missing root IDs
        for root_id in missing_root_ids:
            new_sheet = self._create_sheet_for_material_line(root_id)
            if new_sheet.get('list'):
                # Set initial domain for new list
                new_sheet['list']['domain'] = [['id', 'in', root_to_lines[root_id]]]
            lists[str(root_id)] = new_sheet['list']
            sheets.append(new_sheet['sheet'])

        # Remove sheets
        if removed_ids:
            for rid in removed_ids:
                if str(rid) in lists:
                    del lists[str(rid)]
            sheets = [s for s in sheets if not any(str(rid) in json.dumps(s) for rid in removed_ids)]

        spreadsheet_json['lists'] = lists
        spreadsheet_json['sheets'] = sheets
        
        # 🔍 DEBUG: Log all sheets
        _logger.info(f"📊 Total sheets in spreadsheet: {len(sheets)}")
        for sheet in sheets:
            sheet_id = sheet.get('id', 'unknown')
            sheet_name = sheet.get('name', 'unnamed')
            has_cells = 'cells' in sheet and len(sheet.get('cells', {})) > 0
            cell_count = len(sheet.get('cells', {})) if 'cells' in sheet else 0
            _logger.info(f"   📄 Sheet: {sheet_id} | Name: {sheet_name} | Cells: {cell_count} | Has data: {has_cells}")

        # 🎯 REORDER SHEETS: Main sheets first, then auxiliary sheets
        # Main sheets have IDs like "sheet_123" (material line sheets)
        # Auxiliary sheets have IDs like "profile_master", "resin", "helper"
        main_sheets = []
        auxiliary_sheets = []
        
        for sheet in sheets:
            sheet_id = sheet.get('id', '')
            if sheet_id.startswith('sheet_'):
                # This is a main material line sheet
                main_sheets.append(sheet)
            else:
                # This is an auxiliary sheet (Profile, Resin, Helper, etc.)
                auxiliary_sheets.append(sheet)
        
        # Sort auxiliary sheets by priority
        def sheet_order(s):
            name = (s.get('name') or "").lower()
            if "profile master" in name or "profile" in name:
                return 1
            if "resin" in name:
                return 2
            if "helper" in name:
                return 3
            return 99
        
        auxiliary_sheets.sort(key=sheet_order)
        
        # Rebuild sheets array in correct order
        spreadsheet_json['sheets'] = main_sheets + auxiliary_sheets
        if main_sheets:
            first_main_sheet_id = main_sheets[0].get("id")
            data['active_sheet_id'] = first_main_sheet_id
            spreadsheet_json['active_sheet_id'] = first_main_sheet_id
            
        # ✅ CRITICAL FIX: Preload data for ALL lists
        
        for list_id, list_config in lists.items():
            try:
                line_id = int(list_id)
                line = self.env['crm.material.line'].browse(line_id)
                if line.exists():
                    columns = list_config.get('columns', [])
                    list_data = self.get_list_data('crm.material.line', list_id, columns)
                    
                    # Check if Gel Coat REQ is "No"
                    attrs = line.attributes_json or {}
                    if attrs.get('Gel Coat REQ', '').lower() == 'no':
                        # Find columns to remove
                        cols_to_remove = []
                        for col_idx, col_name in enumerate(columns):
                            # Check if it's a Gel-coat related attribute
                            if 'gel' in col_name.lower() and 'req' not in col_name.lower():
                                cols_to_remove.append(col_idx)
                        
                        # Remove columns in reverse order to maintain indices
                        for col_idx in reversed(cols_to_remove):
                            removed_name = columns.pop(col_idx)
                            
                            # Also remove from data if present
                            if list_data and len(list_data) > 0:
                                for row in list_data:
                                    if removed_name in row:
                                        del row[removed_name]
                            
                            # Check if there's a corresponding UOM column
                            attr_name = removed_name
                            uom_name = f"{attr_name} UOM"
                            
                            # Find and remove UOM column if it exists
                            if uom_name in columns:
                                uom_idx = columns.index(uom_name)
                                columns.pop(uom_idx)
                                if list_data and len(list_data) > 0:
                                    for row in list_data:
                                        if uom_name in row:
                                            del row[uom_name]

                    if list_data:
                        spreadsheet_json['lists'][list_id]['data'] = list_data
            except Exception as e:
                _logger.error(f"❌ Failed to preload list {list_id}: {e}")
        
        # ✅ ENSURE VALIDATION RULES EXIST IN ALL SHEETS
        _logger.info("🔍 Checking validation rules in all sheets...")
        for sheet in spreadsheet_json.get('sheets', []):
            # Check if this is a main sheet (has validations in template)
            sheet_id = sheet.get('id', '')
            _logger.info(f"   🔎 Sheet: {sheet_id} ({sheet.get('name')})")
            _logger.info(f"      Has validations? {'validations' in sheet}")
            _logger.info(f"      Has dataValidationRules? {sheet.get('dataValidationRules') is not None}")
            _logger.info(f"      dataValidationRules count: {len(sheet.get('dataValidationRules', []))}")
            
            if sheet_id.startswith('sheet_'):
                # This is a material line sheet - check if it needs validation rules
                # First, check if sheet already has validations defined
                has_validations_in_sheet = 'validations' in sheet
                has_validation_rules = sheet.get('dataValidationRules') and len(sheet.get('dataValidationRules', [])) > 0
                
                # If no validations in sheet, try to load from template
                validations_to_process = sheet.get('validations', [])
                
                if not has_validations_in_sheet:
                    # Try to get validations from product category template
                    try:
                        line_id = int(sheet_id.replace('sheet_', ''))
                        line = self.env['crm.material.line'].browse(line_id)
                        if line.exists() and line.product_template_id.categ_id.spreadsheet_data:
                            template_data = json.loads(line.product_template_id.categ_id.spreadsheet_data)
                            if template_data and template_data.get('sheets'):
                                main_sheet = template_data['sheets'][0]
                                if 'validations' in main_sheet:
                                    validations_to_process = main_sheet['validations']
                                    _logger.info(f"   📋 Loaded {len(validations_to_process)} validations from template for {sheet.get('name')}")
                    except Exception as e:
                        _logger.warning(f"Failed to load validations from template: {e}")
                
                # Now process validations if we have any and no rules exist yet
                if validations_to_process and not has_validation_rules:
                    _logger.info(f"   ➕ Adding validation rules to sheet {sheet.get('name')}...")
                    sheet['dataValidationRules'] = []
                    for val in validations_to_process:
                        for rng in val.get('ranges', []):
                            try:
                                from openpyxl.utils.cell import range_boundaries, get_column_letter
                                min_col, min_row, max_col, max_row = range_boundaries(rng)
                                for r in range(min_row, max_row + 1):
                                    for c in range(min_col, max_col + 1):
                                        cell_ref = f"{get_column_letter(c)}{r + 4}"  # +4 for header offset
                                        rule_data = self._get_validation_rule(val, [])
                                        if rule_data and rule_data.get('range'):
                                            validation_rule = {
                                                'id': uuid.uuid4().hex,
                                                'isBlocking': False,
                                                'ranges': [cell_ref],
                                                'criterion': {
                                                    'type': 'isValueInRange',
                                                    'values': [rule_data['range']],
                                                    'displayStyle': 'arrow'
                                                }
                                            }
                                            sheet['dataValidationRules'].append(validation_rule)
                                            _logger.info(f"      ✅ Added dropdown at {cell_ref}: {rule_data['range']}")
                            except Exception as e:
                                _logger.warning(f"Failed to add validation: {e}")
                
                # 🔍 DEBUG: Check for formulas that reference other sheets
                if 'cells' in sheet:
                    formula_count = 0
                    for cell_ref, cell_data in sheet['cells'].items():
                        content = cell_data.get('content', '')
                        if isinstance(content, str) and content.startswith('=') and "'" in content:
                            formula_count += 1
                            if formula_count <= 5:  # Show first 5 formulas
                                _logger.info(f"      🔢 Formula in {cell_ref}: {content[:80]}")
                    if formula_count > 0:
                        _logger.info(f"      📐 Total formulas with sheet references: {formula_count}")
        
        data['data'] = spreadsheet_json
        self.raw_spreadsheet_data = json.dumps(spreadsheet_json)

        return data

    # ------------------------------------------------------------------
    # HELPER: Get Columns
    # ------------------------------------------------------------------
    def _get_material_line_columns(self, line):
        """
        Helper to construct columns for a material line sheet.
        Removes 'UOM' and places 'Quantity UOM' next to 'quantity'.
        Also removes Gel-coat column if 'Gel Coat REQ' is 'No'.
        """
        # 1. Base Fields
        columns = list(CRM_MATERIAL_LINE_BASE_FIELDS)

        # 2. Dynamic Attributes
        dynamic_keys = list(line.attributes_json.keys()) if isinstance(line.attributes_json, dict) else []

        # 🔥 Check if Gel Coat REQ is set to "No"
        gel_coat_required = True
        attrs_json = line.attributes_json or {}
        for key, value in attrs_json.items():
            if key and "gel coat req" in key.lower():
                if value and str(value).lower() == "no":
                    gel_coat_required = False
                    
                    break

        # Remove 'UOM' if present (User request: "default UOM... nahi chahiye")
        # We filter it out from dynamic keys AND base fields (just in case)
        if 'UOM' in dynamic_keys:
            dynamic_keys.remove('UOM')
        if 'uom_id' in columns:
            columns.remove('uom_id')

        # Remove 'Quantity' if present (It is already in base fields as 'quantity')
        if 'Quantity' in dynamic_keys:
            dynamic_keys.remove('Quantity')

        # 🔥 Remove Gel-coat related columns if Gel Coat REQ is "No"
        if not gel_coat_required:
            # Check product template attributes to find Gel-coat columns
            template = line.product_template_id
            if template:
                for ptal in template.attribute_line_ids:
                    if ptal.attribute_id.is_gelcoat_required_flag:
                        attr_name = ptal.attribute_id.name
                        # Check if it's NOT the "Gel Coat REQ" attribute itself
                        if attr_name and "gel coat req" not in attr_name.lower():
                            if attr_name in dynamic_keys:
                                
                                dynamic_keys.remove(attr_name)
                            # Also remove UOM variant if exists
                            uom_name = f"{attr_name} UOM"
                            if uom_name in dynamic_keys:
                                
                                dynamic_keys.remove(uom_name)

        # 3. Handle 'Quantity UOM' placement
        qty_uom_key = "Quantity UOM"
        has_qty_uom = False
        if qty_uom_key in dynamic_keys:
            has_qty_uom = True
            dynamic_keys.remove(qty_uom_key)

        # 4. Priority Logic for remaining attributes
        priority = []
        template = line.product_template_id
        if template:
            for ptal in template.attribute_line_ids:
                attr_name = ptal.attribute_id.name
                if attr_name in dynamic_keys:
                    priority.append(attr_name)
                uom_name = f"{attr_name} UOM"
                if uom_name in dynamic_keys:
                    priority.append(uom_name)

        ordered_dynamic = []
        dynamic_keys_copy = list(dynamic_keys)

        for p in priority:
            if p in dynamic_keys_copy:
                ordered_dynamic.append(p)
                dynamic_keys_copy.remove(p)

        ordered_dynamic.extend(sorted(dynamic_keys_copy))

        # 5. Assemble final list
        # Insert Quantity UOM after quantity
        if has_qty_uom:
            if 'quantity' in columns:
                idx = columns.index('quantity') + 1
                columns.insert(idx, qty_uom_key)
            else:
                columns.append(qty_uom_key)

        columns.extend(ordered_dynamic)
        return columns

    def _get_template_commands(self, sheet_id, template_data, row_offset=0):
        """
        Generate commands to recreate the template sheet.
        """
        commands = []
        
        # 1. Cells
        if 'sheets' in template_data and template_data['sheets']:
            template_sheet = template_data['sheets'][0]
            
            # Merges
            from openpyxl.utils.cell import range_boundaries, get_column_letter
            
            for merge in template_sheet.get('merges', []):
                try:
                    # merge is "A1:B2"
                    min_col, min_row, max_col, max_row = range_boundaries(merge)
                    # Shift rows
                    min_row += row_offset
                    max_row += row_offset
                    
                    new_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
                    commands.append({
                        'type': 'ADD_MERGE',
                        'sheetId': sheet_id,
                        'target': [new_range]
                    })
                except Exception:
                    pass
                
            # Columns (No offset needed for columns)
            for col_idx, col_data in template_sheet.get('cols', {}).items():
                commands.append({
                    'type': 'RESIZE_COLUMNS_ROWS',
                    'sheetId': sheet_id,
                    'dimension': 'COL',
                    'elements': [int(col_idx)],
                    'size': col_data.get('width', 100) * 7 
                })
                
            # Rows (Offset needed)
            for row_idx, row_data in template_sheet.get('rows', {}).items():
                commands.append({
                    'type': 'RESIZE_COLUMNS_ROWS',
                    'sheetId': sheet_id,
                    'dimension': 'ROW',
                    'elements': [int(row_idx) + row_offset],
                    'size': row_data.get('size', 21)
                })

            # Cells (Offset needed)
            for cell_ref, cell_data in template_sheet.get('cells', {}).items():
                # Convert A1 to col/row
                col_letter = "".join(filter(str.isalpha, cell_ref))
                row_num = int("".join(filter(str.isdigit, cell_ref))) - 1
                from openpyxl.utils import column_index_from_string
                col_num = column_index_from_string(col_letter) - 1
                
                commands.append({
                    'type': 'UPDATE_CELL',
                    'sheetId': sheet_id,
                    'col': col_num,
                    'row': row_num + row_offset,
                    'content': str(cell_data.get('content', '')) if cell_data.get('content') is not None else '',
                })
                
        return commands

    def _shift_formula_rows(self, content, sheet_name, offset):
        if offset == 0 or not sheet_name:
            return content
        
        escaped_name = re.escape(sheet_name)
        # Match 'Sheet Name'!Ref or SheetName!Ref
        pattern = re.compile(f"(?:'({escaped_name})'|({escaped_name}))!(\\$?[A-Za-z]+)(\\$?)(\\d+)", re.IGNORECASE)
        
        def replace_func(match):
            prefix = match.group(1) or match.group(2)
            is_quoted = bool(match.group(1))
            col_part = match.group(3)
            row_anchor = match.group(4)
            row_num = int(match.group(5))
            new_row = row_num + offset
            sheet_part = f"'{prefix}'" if is_quoted else prefix
            return f"{sheet_part}!{col_part}{row_anchor}{new_row}"

        return pattern.sub(replace_func, content)

    def _fix_formula(self, content, original_name, new_name, all_sheet_names, main_sheet_info=None):
        if not content or not isinstance(content, str) or not content.startswith('='):
            return content
        
        # 0. Fix triple quotes FIRST (before any other processing)
        # This handles """COMPUTED_VALUE""" -> COMPUTED_VALUE
        if '"""' in content:
            # Handle triple quotes - extract content between them
            content = re.sub(r'"""([^"]*?)"""', r'\1', content)
            _logger.info(f"✅ Fixed triple quotes: {content[:100]}")
        
        # 1. Remove __xludf.dummyfunction wrapper - IMPROVED VERSION
        # This wrapper appears when Excel formulas with UNIQUE, FILTER, IFERROR are imported
        if "__xludf.dummyfunction" in content:
            original = content
            
            def unwrap_xludf(formula):
                """
                Iteratively unwrap __xludf.dummyfunction() calls.
                Handles nested quotes and complex formulas.
                """
                max_iterations = 10  # Prevent infinite loop
                iteration = 0
                
                while '__xludf.dummyfunction' in formula and iteration < max_iterations:
                    # Try to extract content - handle nested quotes by matching to the end
                    # Pattern: __xludf.dummyfunction("...content...") where content may have nested quotes
                    match = re.search(r'__xludf\.dummyfunction\("(.+)"\s*\)\s*$', formula, re.DOTALL)
                    
                    if not match:
                        # Try simpler pattern for cases with no nesting
                        match = re.search(r'__xludf\.dummyfunction\("([^"]+)"\)', formula)
                    
                    if match:
                        extracted = match.group(1)
                        # Replace the wrapper with just the extracted content
                        formula = formula[:match.start()] + extracted + formula[match.end():]
                        _logger.info(f"🔧 Unwrapped xludf iteration {iteration + 1}: {extracted[:60]}...")
                    else:
                        # Fallback: simple string replacement
                        # This handles cases where regex fails due to complex nesting
                        if formula.count('__xludf.dummyfunction("') == 1:
                            formula = formula.replace('__xludf.dummyfunction("', '', 1)
                            # Find the last occurrence of ") and remove it
                            if '")' in formula:
                                last_idx = formula.rfind('")')
                                formula = formula[:last_idx] + formula[last_idx+2:]
                        break
                    
                    iteration += 1
                
                return formula
            
            content = unwrap_xludf(content)
            if content != original:
                _logger.info(f"✅ Removed __xludf wrapper: {original[:80]} → {content[:80]}")
        
        # 2. Shift references to Main Sheet if needed
        if main_sheet_info:
            ms_name = main_sheet_info.get('name')
            ms_offset = main_sheet_info.get('offset', 0)
            if ms_name and ms_offset:
                content = self._shift_formula_rows(content, ms_name, ms_offset)

        # 3. Fix self-reference (rename current sheet in formulas)
        if original_name and new_name and original_name != new_name:
            # Replace 'Original Name'! with 'New Name'!
            content = content.replace(f"'{original_name}'!", f"'{new_name}'!")
            if ' ' not in original_name:
                content = content.replace(f"{original_name}!", f"'{new_name}'!")
        
        # 4. Fix references to main sheet in auxiliary sheets
        # When main "Costing" sheet is renamed to product name, update references
        if main_sheet_info and main_sheet_info.get('name') and main_sheet_info.get('new_name'):
            old_main = main_sheet_info['name']
            new_main = main_sheet_info['new_name']
            if old_main != new_main:
                # Update both quoted and unquoted references
                content = content.replace(f"'{old_main}'!", f"'{new_main}'!")
                if ' ' not in old_main:
                    # Use regex to avoid partial matches
                    content = re.sub(
                        f"\\b{re.escape(old_main)}!",
                        f"'{new_main}'!",
                        content
                    )

        # 5. Ensure all sheet names in cross-sheet references are properly quoted
        # This prevents #NAME? errors
        if all_sheet_names:
            for sheet_name in all_sheet_names:
                if not sheet_name:
                    continue
                # If sheet name contains spaces or special chars, ensure it's quoted
                if ' ' in sheet_name or '-' in sheet_name or '.' in sheet_name:
                    # Replace unquoted sheet references with quoted ones
                    # Pattern: SheetName! but not 'SheetName'!
                    escaped_name = re.escape(sheet_name)
                    # Look for unquoted references
                    content = re.sub(
                        f"(?<!')[\\b]{escaped_name}(?=!)",
                        f"'{sheet_name}'",
                        content
                    )
        
        # 6. Fix UNIQUE and FILTER formulas - ensure sheet references are quoted
        if 'UNIQUE' in content or 'FILTER' in content:
            # Ensure proper sheet references in FILTER function
            # Pattern: FILTER(Sheet!Range, Condition)
            def fix_filter_syntax(match):
                sheet_range = match.group(1)
                condition = match.group(2)
                
                # Ensure sheet names are quoted
                if '!' in sheet_range and not sheet_range.startswith("'"):
                    sheet_part, range_part = sheet_range.split('!', 1)
                    sheet_range = f"'{sheet_part}'!{range_part}"
                
                return f"FILTER({sheet_range}, {condition})"
            
            # Apply the fix only if needed
            content = re.sub(r'FILTER\(([^,]+),\s*(.+?)\)', fix_filter_syntax, content)
        
        # 7. Ensure no double equals (prevents #BAD EXPR)
        if content.startswith("=="):
            content = content[1:]
        
        # 8. Fix common Excel → Odoo incompatibilities
        # Remove extra single quotes that might cause issues
        content = content.replace("''", "'")
            
        return content

    def _get_validation_rule(self, val, all_sheet_names):
        """
        Convert Excel dropdown formula into correct Odoo Spreadsheet dataValidation rule.
        Supports:
            • "A,B,C" static list
            • =Sheet!$A$2:$A$1000 range dropdown
        """
        f1 = str(val.get('formula1', '')).strip()
        

        # -------- STATIC LIST -------- ("A,B,C")
        if (f1.startswith('"') and f1.endswith('"')) or ("," in f1 and not f1.startswith("=")):
            raw = f1.strip('"')
            values = [v.strip() for v in raw.split(',')]
            
            return {'type': 'range', 'values': values, 'style': 'arrow'}

        # -------- RANGE DROPDOWN -------- (=Sheet!$A$2:$A$100)
        if f1.startswith("="):
            f1 = f1[1:]  # remove "=" temporarily for cleaning

        # ✅ FIX: Do NOT remove '$'. Keep absolute references for dropdowns.
        clean = f1.strip().strip('"')  

        # auto-quote ALWAYS (Odoo prefers quoted sheet names for validation ranges)
        if "!" in clean:
            sheet_part, range_part = clean.split("!", 1)
            # Strip existing quotes to avoid double quoting
            sheet_part = sheet_part.strip("'")
            clean = f"'{sheet_part}'!{range_part}"

        
        return {'type': 'range', 'range': clean, 'style': 'arrow'}


    def _get_sheet_populate_commands(
        self, sheet_json, sheet_id, row_offset=0, new_sheet_name=None,
        all_sheet_names=None, main_sheet_info=None
    ):
        """
        Populate sheet with merge, resize, formula-shift and Excel → Odoo dropdown conversion.
        """
        commands = []
        original_sheet_name_raw = sheet_json.get('name') or ""
        original_sheet_name = original_sheet_name_raw.lower()

        # 🔥 Detect reference sheets (No row offset for helper, resin, profile master etc.)
        reference_keywords = ["helper", "resin", "profile", "profile master", "master"]
        if any(k in original_sheet_name for k in reference_keywords):
            row_offset = 0
            col_offset = 0
        else:
            row_offset = 4
            col_offset = 0

        # -----------------------------------------
        # MERGES
        # -----------------------------------------
        for merge in sheet_json.get('merges', []):
            try:
                min_col, min_row, max_col, max_row = range_boundaries(merge)
                min_row += row_offset; max_row += row_offset
                min_col += col_offset; max_col += col_offset
                new_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
                commands.append({'type': 'ADD_MERGE', 'sheetId': sheet_id, 'target': [new_range]})
            except:
                pass

        # -----------------------------------------
        # COLUMN WIDTHS
        # -----------------------------------------
        for col_idx, col_data in sheet_json.get('cols', {}).items():
            commands.append({
                'type': 'RESIZE_COLUMNS_ROWS', 'sheetId': sheet_id,
                'dimension': 'COL', 'elements': [int(col_idx)],
                'size': col_data.get('width', 100) * 7
            })

        # -----------------------------------------
        # ROW HEIGHTS
        # -----------------------------------------
        for row_idx, row_data in sheet_json.get('rows', {}).items():
            commands.append({
                'type': 'RESIZE_COLUMNS_ROWS', 'sheetId': sheet_id,
                'dimension': 'ROW', 'elements': [int(row_idx) + row_offset],
                'size': row_data.get('size', 21)
            })

        # -----------------------------------------
        # Build dropdown validation map (Excel → Odoo conversion)
        # -----------------------------------------
        cell_validations = {}
        if 'validations' in sheet_json:
            _logger.info(f"🔍 [_get_sheet_populate_commands] Found {len(sheet_json['validations'])} validations in sheet")
            for val in sheet_json['validations']:
                for rng in val.get('ranges', []):
                    try:
                        min_col, min_row, max_col, max_row = range_boundaries(rng)
                        for r in range(min_row, max_row + 1):
                            for c in range(min_col, max_col + 1):
                                cell_validations[(r - 1, c - 1)] = self._get_validation_rule(val, all_sheet_names)
                    except:
                        pass
            _logger.info(f"   📋 Built {len(cell_validations)} cell validation mappings")
        else:
            _logger.warning(f"   ❌ No 'validations' key in sheet_json")

        # -----------------------------------------
        # CELLS & FORMULAS
        # -----------------------------------------
        processed_cells = set()
        
        for cell_ref, cell_data in sheet_json.get('cells', {}).items():
            col_l = "".join(filter(str.isalpha, cell_ref))
            row_n = int("".join(filter(str.isdigit, cell_ref))) - 1
            col_n = column_index_from_string(col_l) - 1
            
            processed_cells.add((row_n, col_n))

            content = str(cell_data.get('content', '')) if cell_data.get('content') else ''

            # 🔁 Relocate formulas to new shifted position
            if content.startswith('='):
                original_formula = content
                
                # 🔥 CRITICAL FIX: For formulas with cross-sheet references, 
                # we DON'T use Translator because it breaks them
                # Instead, we manually shift only SAME-SHEET references
                
                has_cross_sheet = "!" in content
                
                if has_cross_sheet:
                    # Manual shift for cross-sheet formulas
                    # Only shift references WITHOUT sheet names (same-sheet refs)
                    
                    def shift_same_sheet_ref(match):
                        """Shift only references without sheet prefix"""
                        full_match = match.group(0)
                        col_abs = match.group(1) or ''
                        col = match.group(2)
                        row_abs = match.group(3) or ''
                        row = int(match.group(4))
                        
                        # Check if this is part of a sheet reference (preceded by ! or ')
                        start_pos = match.start()
                        if start_pos > 0:
                            prev_char = content[start_pos - 1]
                            if prev_char in ['!', "'"]:
                                # This is a cross-sheet reference, don't shift
                                return full_match
                        
                        # This is a same-sheet reference, apply offset
                        if row_abs == '$':
                            # Absolute row reference
                            new_row = row + row_offset
                        else:
                            # Relative row reference  
                            new_row = row + row_offset
                        
                        return f"{col_abs}{col}{row_abs}{new_row}"
                    
                    # Pattern to match cell references (but not those after ! or ')
                    pattern = r'(\$?)([A-Z]+)(\$?)(\d+)'
                    content = re.sub(pattern, shift_same_sheet_ref, content)
                    
                else:
                    # No cross-sheet refs, use Translator (original logic)
                    try:
                        origin = f"{col_l}{row_n + 1}"
                        dest_col_letter = get_column_letter(col_n + col_offset + 1)
                        dest_ref = f"{dest_col_letter}{row_n + row_offset + 1}"
                        content = Translator(content, origin=origin).translate_formula(dest_ref)
                    except Exception as e:
                        _logger.warning(f"Translator failed for {cell_ref}: {e}")
                
                content = self._fix_formula(content, original_sheet_name_raw, new_sheet_name, all_sheet_names, main_sheet_info)

            cmd = {
                'type': 'UPDATE_CELL', 'sheetId': sheet_id,
                'col': col_n + col_offset, 'row': row_n + row_offset,
                'content': content,
            }

            commands.append(cmd)

        # -----------------------------------------
        # 🔥 STEP 3: Apply ALL validations as SEPARATE commands
        # -----------------------------------------
        _logger.info(f"🎯 Applying {len(cell_validations)} validations as separate commands...")
        
        validation_commands = []
        for (r, c), rule in cell_validations.items():
            target_col = c + col_offset
            target_row = r + row_offset
            cell_range = f"{get_column_letter(target_col + 1)}{target_row + 1}"
            
            # ✅ Build Odoo-compatible validation structure
            validation = None
            
            if rule.get('range'):
                # 🔥 Range-based dropdown (e.g., 'Profile Master'!A2:A1000)
                src = rule['range']
                
                # Clean range reference
                if src.startswith('"') and src.endswith('"'):
                    src = src[1:-1]
                if src.startswith("''") and not src.startswith("'''"):
                    src = src[1:]
                
                validation = {
                    'type': 'isValueInRange',
                    'values': [src],
                    'displayStyle': 'arrow'
                }
                _logger.info(f"✅ Range validation @ {cell_range}: {src}")
                
            elif rule.get('values'):
                # 🔥 Static list dropdown (e.g., ["A", "B", "C"])
                validation = {
                    'type': 'isValueInList',
                    'values': rule['values'],
                    'displayStyle': 'arrow'
                }
                _logger.info(f"✅ List validation @ {cell_range}: {rule['values']}")
            
            if validation:
                dv_id = f"dv_{uuid.uuid4().hex[:8]}"
                
                validation_cmd = {
                    'type': 'ADD_DATA_VALIDATION_RULE',
                    'sheetId': sheet_id,
                    'ranges': [cell_range],
                    'rule': {
                        'id': dv_id,
                        'criterion': validation
                    }
                }
                validation_commands.append(validation_cmd)
        
        commands.extend(validation_commands)
        _logger.info(f"🎯 Added {len(validation_commands)} validation commands to dispatch queue")

        return commands

    # ------------------------------------------------------------------
    # EMPTY DATA (initial load)
    # ------------------------------------------------------------------
    def _empty_spreadsheet_data(self):
        data = super()._empty_spreadsheet_data() or {}
        data.setdefault('lists', {})
        data['sheets'] = []

        if not self.lead_id or not self.lead_id.material_line_ids:
            return data

        # ✅ Group lines by root for collective sheet initialization
        root_to_lines = {}
        for line in self.lead_id.material_line_ids:
            _, root_id_str = self._get_line_root_info(line)
            root_id = int(root_id_str)
            if root_id not in root_to_lines:
                root_to_lines[root_id] = []
            root_to_lines[root_id].append(line)

        for root_id, lines in root_to_lines.items():
            first_line = lines[0]
            sheet_id, list_id = self._get_line_root_info(first_line)
            product_name = (first_line.product_template_id.display_name or "Item")[:31]

            columns = self._get_material_line_columns(first_line)

            # Check for template
            template_data = None
            
            # Force reload category to bypass cache
            category = first_line.product_template_id.categ_id
            category.invalidate_recordset(['spreadsheet_data'])
            
            if category.spreadsheet_data:
                try:
                    _logger.info(f"🔍 Loading template for category: {category.name}")
                    template_data = json.loads(category.spreadsheet_data)
                    if template_data and template_data.get('sheets'):
                        t_sheet = template_data['sheets'][0]
                        _logger.info(f"   � Template Sheet Keys: {list(t_sheet.keys())}")
                        if 'validations' in t_sheet:
                            _logger.info(f"   ✅ Template has {len(t_sheet['validations'])} validations")
                        else:
                            _logger.warning("   ❌ Template has NO 'validations' key")
                except Exception as e:
                    _logger.error(f"Failed to load template data: {e}")
                    pass

            if template_data and template_data.get('sheets'):
                # Use template sheet
                template_sheet = template_data['sheets'][0]
                # We need to deep copy to avoid modifying the cached template
                import copy
                sheet_json = copy.deepcopy(template_sheet)
                sheet_json['id'] = sheet_id
                sheet_json['name'] = product_name
                
                # ✅ SANITIZE: Ensure all cell content is string AND remove invalid format
                if 'cells' in sheet_json:
                    for cell_key, cell_val in sheet_json['cells'].items():
                        if 'content' in cell_val:
                            cell_val['content'] = str(cell_val['content'])
                        if 'format' in cell_val:
                             # Remove format to prevent #ERROR if format ID is invalid in new sheet
                            del cell_val['format']

                # ✅ Add dataValidationRules from template validations
                if 'validations' in sheet_json and not sheet_json.get('dataValidationRules'):
                    _logger.info(f"🔍 [_empty_spreadsheet_data] Adding validations to MAIN sheet {sheet_json.get('name')}")
                    sheet_json['dataValidationRules'] = []
                    for val in sheet_json['validations']:
                        for rng in val.get('ranges', []):
                            try:
                                from openpyxl.utils.cell import range_boundaries
                                min_col, min_row, max_col, max_row = range_boundaries(rng)
                                # Apply row offset of 4 for main sheet
                                for r in range(min_row, max_row + 1):
                                    for c in range(min_col, max_col + 1):
                                        cell_ref = f"{get_column_letter(c)}{r + 4}"  # +4 for header offset
                                        rule_data = self._get_validation_rule(val, [])
                                        if rule_data:
                                            validation_rule = {
                                                'id': uuid.uuid4().hex,
                                                'isBlocking': False,
                                                'ranges': [cell_ref],
                                            }
                                            if rule_data.get('range'):
                                                validation_rule['criterion'] = {
                                                    'type': 'isValueInRange',
                                                    'values': [rule_data['range']],
                                                    'displayStyle': 'arrow'
                                                }
                                            elif rule_data.get('values'):
                                                validation_rule['criterion'] = {
                                                    'type': 'isValueInList',
                                                    'values': rule_data['values'],
                                                    'displayStyle': 'arrow'
                                                }
                                            if validation_rule.get('criterion'):
                                                sheet_json['dataValidationRules'].append(validation_rule)
                                                _logger.info(f"✅ [_empty_spreadsheet_data] Added validation to {cell_ref}: {rule_data.get('range') or rule_data.get('values')}")
                            except Exception as e:
                                _logger.warning(f"Failed to process validation for {rng}: {e}")

                _logger.info(f"📋 Sheet {product_name}: validations={('validations' in sheet_json)}, dataValidationRules={len(sheet_json.get('dataValidationRules', []))}")
                main_sheets.append(sheet_json)

                # 2. Auxiliary Sheets (Index 1+) - Collect for sorting
                if len(template_data['sheets']) > 1:
                    for aux_sheet in template_data['sheets'][1:]:
                        # Check if sheet with this ID already exists
                        existing_ids = {s.get('id') for s in auxiliary_sheets}
                        if aux_sheet.get('id') not in existing_ids:
                            aux_copy = copy.deepcopy(aux_sheet)
                            # Sanitize
                            if 'cells' in aux_copy:
                                for cell_key, cell_val in aux_copy['cells'].items():
                                    if 'content' in cell_val:
                                        cell_val['content'] = str(cell_val['content']) if cell_val['content'] is not None else ""
                                    if 'format' in cell_val:
                                        del cell_val['format']
                            
                            # ✅ Add dataValidationRules from template validations
                            
                            # ✅ Add dataValidationRules from template validations
                            if 'validations' in aux_copy and not aux_copy.get('dataValidationRules'):
                                aux_copy['dataValidationRules'] = []
                                for val in aux_copy['validations']:
                                    for rng in val.get('ranges', []):
                                        try:
                                            from openpyxl.utils.cell import range_boundaries
                                            min_col, min_row, max_col, max_row = range_boundaries(rng)
                                            for r in range(min_row, max_row + 1):
                                                for c in range(min_col, max_col + 1):
                                                    cell_ref = f"{get_column_letter(c)}{r}"
                                                    rule_data = self._get_validation_rule(val, [])
                                                    if rule_data:
                                                        validation_rule = {
                                                            'id': uuid.uuid4().hex,
                                                            'isBlocking': False,
                                                            'ranges': [cell_ref],
                                                        }
                                                        if rule_data.get('range'):
                                                            validation_rule['criterion'] = {
                                                                'type': 'isValueInRange',
                                                                'values': [rule_data['range']],
                                                                'displayStyle': 'arrow'
                                                            }
                                                        elif rule_data.get('values'):
                                                            validation_rule['criterion'] = {
                                                                'type': 'isValueInList',
                                                                'values': rule_data['values'],
                                                                'displayStyle': 'arrow'
                                                            }
                                                        if validation_rule.get('criterion'):
                                                            aux_copy['dataValidationRules'].append(validation_rule)
                                                            _logger.info(f"✅ Added validation to AUX sheet {aux_copy.get('name')}: {cell_ref} -> {rule_data.get('range')}")
                                        except Exception as e:
                                                            _logger.warning(f"Failed to process validation for {rng}: {e}")
                            
                            auxiliary_sheets.append(aux_copy)
            else:
                # Default empty sheet
                main_sheets.append({'id': sheet_id, 'name': product_name})

            data['lists'][list_id] = {
                'id': list_id,
                'model': 'crm.material.line',
                'columns': columns,
                'domain': [['id', 'in', [l.id for l in lines]]],
                'sheetId': sheet_id,
                'name': product_name,
                'context': {},
                'orderBy': [],
                'fieldMatching': {
                    'material_line_ids': {'chain': 'lead_id', 'type': 'many2one'},
                },
            }
        # 🎯 SORT auxiliary sheets: Profile Master → Resin → Helper
        def sheet_order(s):
            name = (s.get('name') or "").lower()
            if "profile master" in name or "profile" in name:
                return 1
            if "resin" in name:
                return 2
            if "helper" in name:
                return 3
            return 99

        auxiliary_sheets.sort(key=sheet_order)

        # ✅ Append in correct order: Main sheets first, then sorted auxiliaries
        data['sheets'].extend(main_sheets)
        data['sheets'].extend(auxiliary_sheets)

        return data

    # ------------------------------------------------------------------
    # INSERT REVISION (create new sheet on new line)
    # ------------------------------------------------------------------
    def _dispatch_insert_list_revision(self):
        self.ensure_one()

        line_id = self._context.get('material_line_id')
        if not line_id:
            return
        
        commands = []
        line = self.env['crm.material.line'].browse(line_id)
        if not line.exists():
            return

        sheet_id, list_id = self._get_line_root_info(line)
        product_name = (line.product_template_id.display_name or "Item")[:31]
        
        # ✅ Load current spreadsheet data to check for existence
        current_data = json.loads(self.raw_spreadsheet_data) if self.raw_spreadsheet_data else {}
        existing_sheet = any(s.get('id') == sheet_id for s in current_data.get('sheets', []))

        # ✅ Get all lines in this specific group (Root + Duplicates)
        group_lines = self.env['crm.material.line'].search([
            '|', ('id', '=', int(list_id)), ('original_line_id', '=', int(list_id))
        ], order='id asc')
        
        domain = [['id', 'in', group_lines.ids]]
        lines_count = len(group_lines)
        first_line = group_lines[0]
        columns = self._get_material_line_columns(first_line)

        # Build column metadata
        columns_meta = []
        for col in columns:
            ftype = self.env['crm.material.line']._fields[col].type if col in self.env['crm.material.line']._fields else 'char'
            columns_meta.append({'name': col, 'type': ftype})

        # ✅ PHASE 1: Sheet & List Setup
        if not existing_sheet:
            # Create NEW sheet
            commands.append({
                'type': 'CREATE_SHEET', 
                'sheetId': sheet_id, 
                'name': product_name,
                'position': 0 
            })
        
        # Always REGISTER or UPDATE the list to include the new domain
        commands.append({
            'type': 'REGISTER_ODOO_LIST',
            'listId': list_id,
            'model': 'crm.material.line',
            'columns': columns,
            'domain': domain,
            'context': {},
            'orderBy': [],
        })
        
        # Always RE_INSERT to update headers and lines
        commands.append({
            'type': 'RE_INSERT_ODOO_LIST',
            'sheetId': sheet_id,
            'col': 0,
            'row': 0,
            'id': list_id,
            'linesNumber': lines_count,
            'columns': columns_meta,
        })

        # ✅ PHASE 2: Data Population (Loop through all lines in the group)
        for row_idx, g_line in enumerate(group_lines, start=1):
            attrs = g_line.attributes_json or {}
            
            for col_idx, col_meta in enumerate(columns_meta):
                field_name = col_meta['name']
                if field_name in g_line._fields:
                    val = g_line[field_name]
                    cell_value = val.display_name if hasattr(val, 'display_name') else (val if val is not False else '')
                else:
                    cell_value = attrs.get(field_name, '')

                commands.append({
                    'type': 'UPDATE_CELL',
                    'sheetId': sheet_id,
                    'col': col_idx,
                    'row': row_idx, # Row 1, 2, 3...
                    'content': str(cell_value) if cell_value not in (None, False, '') else '',
                })

                # Header cleanup (only for the first line/overall header)
                if row_idx == 1 and "__" in field_name:
                    display_name = field_name.split("__")[0]
                    commands.append({
                        'type': 'UPDATE_CELL',
                        'sheetId': sheet_id,
                        'col': col_idx, 'row': 0,
                        'content': display_name,
                    })

        # ✅ PHASE 3: Formatting & Template (Only for fixed categories)
        if not existing_sheet:
            commands.append({
                'type': 'CREATE_TABLE',
                'sheetId': sheet_id,
                'tableType': 'static',
                'ranges': [{
                    '_sheetId': sheet_id,
                    '_zone': {'top': 5, 'bottom': 500, 'left': 1, 'right': len(columns_meta)},
                }],
                'config': {'bandedRows': True, 'styleId': 'TableStyleMedium5', 'hasFilters': True},
            })

        # Check for template and append if exists
        template_data = None
        category = line.product_template_id.categ_id
        _logger.info(f"🔍 [_dispatch_insert_list_revision] Category: {category.name}, Has data: {bool(category.spreadsheet_data)}")
        
        if category.spreadsheet_data:
            try:
                template_data = json.loads(category.spreadsheet_data)
                if template_data and template_data.get('sheets'):
                    main_sheet = template_data['sheets'][0]
                    _logger.info(f"   📋 Main sheet keys: {list(main_sheet.keys())}")
                    if 'validations' in main_sheet:
                        _logger.info(f"   ✅ Main sheet has {len(main_sheet['validations'])} validations")
                    else:
                        _logger.warning(f"   ❌ Main sheet has NO validations key")
            except Exception as e:
                _logger.error(f"   ❌ Failed to parse template: {e}")
                pass

        if template_data and template_data.get('sheets'):
            # Collect all sheet names for dynamic reference fixing
            all_sheet_names = [s.get('name') for s in template_data.get('sheets', [])]

            # 🎯 SORT auxiliary sheets: Profile Master → Resin → Helper → Others
            # This ensures correct tab order when spreadsheet opens
            aux_sheets = template_data['sheets'][1:]
            
            def sheet_order(s):
                name = (s.get('name') or "").lower()
                if "profile master" in name or "profile" in name:
                    return 1
                if "resin" in name:
                    return 2
                if "helper" in name:
                    return 3
                return 99

            aux_sheets.sort(key=sheet_order)

            # ---------------------------------------------------------
            # PREPARE: Main Sheet Info (Needed for Aux sheets too)
            # ---------------------------------------------------------
            main_sheet = template_data['sheets'][0]
            main_sheet_name = main_sheet.get('name')
            main_sheet_info = {
                'name': main_sheet_name, 
                'new_name': product_name,  # ✅ New name after merge
                'offset': 4
            }

            # ---------------------------------------------------------
            # PHASE 1: Create Auxiliary Sheets FIRST
            # (Ensures cross-sheet validations work when Main Sheet is populated)
            # ---------------------------------------------------------
            aux_create_cmds = []
            aux_populate_cmds = []
            
            current_data = json.loads(self.raw_spreadsheet_data) if self.raw_spreadsheet_data else {}
            current_sheet_ids = {s.get('id') for s in current_data.get('sheets', [])}

            for idx, aux_sheet in enumerate(aux_sheets, start=1):
                aux_id = aux_sheet.get('id')
                
                # Check if already exists in current data OR in pending revisions
                is_duplicate = False
                if aux_id in current_sheet_ids:
                    is_duplicate = True
                else:
                    # Check DB for pending commands creating this sheet
                    domain = [
                        ('res_id', '=', self.id),
                        ('res_model', '=', self._name),
                        ('commands', 'ilike', aux_id)
                    ]
                    if self.env['spreadsheet.revision'].search_count(domain):
                        is_duplicate = True

                if aux_id and not is_duplicate:
                    # 1. Create Command
                    aux_create_cmds.append({
                        'type': 'CREATE_SHEET', 
                        'sheetId': aux_id, 
                        'name': aux_sheet.get('name', 'Sheet')[:31],
                        'position': idx  # ✅ Force position 1, 2, 3...
                    })
                    
                    # 2. Populate Command (Stored for later)
                    # Auxiliary sheets (Helper, Resin, Profile Master) are reference sheets
                    # They should NOT have row offset applied
                    p_cmds = self._get_sheet_populate_commands(
                        aux_sheet, 
                        aux_id, 
                        row_offset=0,  # ✅ Changed from 4 to 0
                        new_sheet_name=aux_sheet.get('name'),
                        all_sheet_names=all_sheet_names,
                        main_sheet_info=main_sheet_info
                    )
                    aux_populate_cmds.extend(p_cmds)
                    
                    # ✅ ADD VALIDATION RULES to auxiliary sheet
                    if 'validations' in aux_sheet:
                        _logger.info(f"🔍 Adding {len(aux_sheet['validations'])} validation rules to AUX sheet {aux_sheet.get('name')}")
                        for val in aux_sheet['validations']:
                            for rng in val.get('ranges', []):
                                try:
                                    from openpyxl.utils.cell import range_boundaries
                                    min_col, min_row, max_col, max_row = range_boundaries(rng)
                                    for r in range(min_row, max_row + 1):
                                        for c in range(min_col, max_col + 1):
                                            cell_ref = f"{get_column_letter(c)}{r}"
                                            rule_data = self._get_validation_rule(val, [])
                                            if rule_data and rule_data.get('range'):
                                                # Create UPDATE_CELL command with dataValidation
                                                validation_cmd = {
                                                    'type': 'UPDATE_CELL',
                                                    'sheetId': aux_id,
                                                    'col': c - 1,
                                                    'row': r - 1,
                                                    'dataValidation': {
                                                        'id': uuid.uuid4().hex,
                                                        'criterion': {
                                                            'type': 'isValueInRange',
                                                            'values': [rule_data['range']],
                                                            'displayStyle': 'arrow'
                                                        },
                                                        'ranges': [cell_ref]
                                                    }
                                                }
                                                aux_populate_cmds.append(validation_cmd)
                                                _logger.info(f"✅ Added dropdown to AUX {aux_sheet.get('name')} at {cell_ref}: {rule_data['range']}")
                                except Exception as e:
                                    _logger.warning(f"Failed to add validation for {rng}: {e}")
                    
                    # Mark as added to avoid duplicates
                    current_sheet_ids.add(aux_id)

            # ✅ Execute Creation Commands FIRST
            commands.extend(aux_create_cmds)

            # ---------------------------------------------------------
            # PHASE 2: Populate Auxiliary Sheets
            # (Populate these FIRST so data exists for Main Sheet validations)
            # ---------------------------------------------------------
            commands.extend(aux_populate_cmds)

            # ---------------------------------------------------------
            # PHASE 3: Populate Main Sheet
            # ---------------------------------------------------------
            # Pass new_sheet_name=product_name to fix self-references
            template_cmds = self._get_sheet_populate_commands(
                main_sheet, 
                sheet_id, 
                row_offset=4, 
                new_sheet_name=product_name,
                all_sheet_names=all_sheet_names,
                main_sheet_info=main_sheet_info
            )
            commands.extend(template_cmds)
            

        # Final update command
        commands.append({'type': 'UPDATE_ODOO_LIST_DATA', 'listId': list_id})

        # Dispatch all commands
        self._dispatch_commands(commands)
        
        # 🔥 CRITICAL FIX: Directly add dataValidationRules to raw_spreadsheet_data
        # Commands alone don't persist validations - we must update the stored JSON
        # IMPORTANT: Reload data AFTER dispatch to get the latest state
        _logger.info(f"🔍 [DIRECT UPDATE] Checking if we should add validations directly...")
        _logger.info(f"   template_data exists: {bool(template_data)}")
        if template_data:
            _logger.info(f"   template_data has sheets: {bool(template_data.get('sheets'))}")
        
        if template_data and template_data.get('sheets'):
            main_sheet = template_data['sheets'][0]
            _logger.info(f"   main_sheet has validations: {'validations' in main_sheet}")
            
            if 'validations' in main_sheet:
                try:
                    # Reload data to get the state AFTER commands were processed
                    self.invalidate_recordset(['raw_spreadsheet_data'])
                    data = json.loads(self.raw_spreadsheet_data) if self.raw_spreadsheet_data else {}
                    sheets = data.get('sheets', [])
                    _logger.info(f"   Current spreadsheet has {len(sheets)} sheets (after reload)")
                    
                    # Find the sheet we just created
                    target_sheet = None
                    for sheet in sheets:
                        if sheet.get('id') == sheet_id:
                            target_sheet = sheet
                            break
                    
                    if target_sheet:
                        _logger.info(f"   ✅ Found target sheet: {sheet_id}")
                        if 'dataValidationRules' not in target_sheet:
                            target_sheet['dataValidationRules'] = []
                        
                        # Add validation rules with proper row offset
                        for val in main_sheet['validations']:
                            for rng in val.get('ranges', []):
                                try:
                                    from openpyxl.utils.cell import range_boundaries
                                    min_col, min_row, max_col, max_row = range_boundaries(rng)
                                    for r in range(min_row, max_row + 1):
                                        for c in range(min_col, max_col + 1):
                                            # Apply row offset of 4
                                            cell_ref = f"{get_column_letter(c)}{r + 4}"
                                            rule_data = self._get_validation_rule(val, all_sheet_names)
                                            
                                            if rule_data and rule_data.get('range'):
                                                validation_rule = {
                                                    'id': uuid.uuid4().hex,
                                                    'criterion': {
                                                        'type': 'isValueInRange',
                                                        'values': [rule_data['range']],
                                                    },
                                                    'ranges': [cell_ref],
                                                    'isBlocking': False
                                                }
                                                target_sheet['dataValidationRules'].append(validation_rule)
                                                _logger.info(f"   ✅ [DIRECT] Added validation rule to {cell_ref}")
                                except Exception as e:
                                    _logger.warning(f"Failed to add validation for {rng}: {e}")
                        
                        # Save updated data
                        self.raw_spreadsheet_data = json.dumps(data)
                        _logger.info(f"   💾 Saved {len(target_sheet.get('dataValidationRules', []))} validation rules to raw_spreadsheet_data")
                    else:
                        _logger.warning(f"   ❌ Could not find sheet {sheet_id} in raw_spreadsheet_data")
                        _logger.info(f"   Available sheet IDs: {[s.get('id') for s in sheets]}")
                except Exception as e:
                    _logger.error(f"Failed to update raw_spreadsheet_data with validations: {e}", exc_info=True)

    # ------------------------------------------------------------------
    # SYNC WITH MATERIAL LINES
    # ------------------------------------------------------------------
    def _sync_sheets_with_material_lines(self):
        self.ensure_one()
        if not self.lead_id:
            return

        data = json.loads(self.raw_spreadsheet_data) if self.raw_spreadsheet_data else {}
        current_sheets = data.get('sheets', [])
        current_lists = data.get('lists', {})
        current_line_ids = set(self.lead_id.material_line_ids.ids)

        # Remove deleted
        for sheet in current_sheets:
            sid = sheet.get('id')
            if sid and sid.startswith("sheet_"):
                try:
                    line_id = int(sid.replace("sheet_", ""))
                    if line_id not in current_line_ids:
                        self._delete_sheet_for_material_line(line_id)
                except Exception:
                    pass

        # Re-add missing OR update if columns changed
        existing_sheet_ids = {
            int(s['id'].replace('sheet_', ''))
            for s in current_sheets
            if s.get('id', '').startswith('sheet_')
        }

        # ✅ Group lines by root ID for correct multi-row sheet assignment
        root_to_lines = {}
        for line in self.lead_id.material_line_ids:
            _, root_id_str = self._get_line_root_info(line)
            root_id = int(root_id_str)
            if root_id not in root_to_lines:
                root_to_lines[root_id] = []
            root_to_lines[root_id].append(line)

        for root_id, parent_lines in root_to_lines.items():
            parent_line = parent_lines[0]
            # Find children
            children = self.env['crm.material.line'].sudo().search([
                ('original_line_id', '=', root_id),
                ('is_dimension_row', '=', True)
            ], order='id asc')
            
            # 2. Check if sheet exists for this root
            if root_id in existing_sheet_ids:
                list_id = str(root_id)
                if list_id in current_lists:
                    current_columns = current_lists[list_id].get('columns', [])
                    # Use a child line for column template if possible
                    sample_line = children[0] if children else parent_line
                    expected_columns = self._get_material_line_columns(sample_line)

                    if current_columns != expected_columns:
                        _logger.info(f"♻️ Columns changed for root {root_id}. Re-creating sheet.")
                        self._delete_sheet_for_material_line(root_id)
                        self.with_context(material_line_id=parent_line.id)._dispatch_insert_list_revision()
                        continue
            
            # 3. Create if missing
            if root_id not in existing_sheet_ids:
                self.with_context(material_line_id=parent_line.id)._dispatch_insert_list_revision()

    # ------------------------------------------------------------------
    # CREATE SHEET STRUCTURE
    # ------------------------------------------------------------------
    def _create_sheet_for_material_line(self, material_line_id):
        self.ensure_one()

        line = self.env['crm.material.line'].browse(material_line_id)
        if not line.exists():
            return {'sheet': {}, 'list': {}}

        sheet_id = f"sheet_{line.id}"
        list_id = str(line.id)
        name = (line.product_template_id.display_name or "Item")[:31]

        columns = self._get_material_line_columns(line)

        return {
            'sheet': {'id': sheet_id, 'name': name},
            'list': {
                'id': list_id,
                'model': 'crm.material.line',
                'columns': columns,
                'domain': [['id', '=', line.id]],
                'sheetId': sheet_id,
                'name': name,
                'context': {},
                'orderBy': [],
                'fieldMatching': {
                    'material_line_ids': {'chain': 'lead_id', 'type': 'many2one'},
                },
            },
        }

    # ------------------------------------------------------------------
    # DELETE SHEET
    # ------------------------------------------------------------------
    def _delete_sheet_for_material_line(self, material_line_id):
        sheet_id = f"sheet_{material_line_id}"
        list_id = str(material_line_id)

        commands = [
            {'type': 'DELETE_SHEET', 'sheetId': sheet_id},
            {'type': 'UNREGISTER_ODOO_LIST', 'listId': list_id},
        ]

        try:
            self._dispatch_commands(commands)
        except:
            self._cleanup_deleted_sheets_from_data(material_line_id)

    def _cleanup_deleted_sheets_from_data(self, material_line_id):
        if not self.raw_spreadsheet_data:
            return
        try:
            data = json.loads(self.raw_spreadsheet_data)
            sid = f"sheet_{material_line_id}"
            if 'sheets' in data:
                data['sheets'] = [s for s in data['sheets'] if s.get('id') != sid]
            if 'lists' in data and str(material_line_id) in data['lists']:
                del data['lists'][str(material_line_id)]
            self.raw_spreadsheet_data = json.dumps(data)
        except:
            pass

    # ------------------------------------------------------------------
    # MANUAL SYNC BUTTON
    # ------------------------------------------------------------------
    def action_sync_sheets(self):
        for spreadsheet in self:
            spreadsheet._sync_sheets_with_material_lines()
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'type': 'success',
                'message': _('Sheets synced with CRM Material Lines'),
                'next': {'type': 'ir.actions.act_window_close'},
            }
        }

    # ------------------------------------------------------------------
    # SPREADSHEET SELECTOR
    # ------------------------------------------------------------------
    @api.model
    def _get_spreadsheet_selector(self):
        return {
            'model': self._name,
            'display_name': _("CRM Quote Spreadsheets"),
            'sequence': 20,
            'allow_create': False,
        }

    # ------------------------------------------------------------------
    # DATA PROVIDER — used by spreadsheets
    # ------------------------------------------------------------------
    @api.model
    def get_crm_material_lines(self):
        self.ensure_one()
        if not self.lead_id:
            return []

        rows = []
        for line in self.lead_id.material_line_ids:

            row = {
                'id': line.id,
                'name': line.product_template_id.display_name if line.product_template_id else '',
                'quantity': line.quantity,
            }

            # Add dynamic JSON attributes
            attrs = line.attributes_json or {}
            for k, v in attrs.items():
                row[k] = v

            rows.append(row)

        return rows

    def getMainCrmMaterialLineLists(self):
        self.ensure_one()
        if not self.lead_id or not self.lead_id.material_line_ids:
            return []

        lists = []
        for line in self.lead_id.material_line_ids:
            columns = self._get_material_line_columns(line)

            lists.append({
                'id': str(line.id),
                'model': 'crm.material.line',
                'field_names': columns,
                'columns': columns,
                'name': line.product_template_id.display_name or f"Item {line.id}",
                'sheetId': f"sheet_{line.id}",
            })

        return lists