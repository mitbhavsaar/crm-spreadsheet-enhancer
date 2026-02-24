# -*- coding: utf-8 -*-
from odoo import models, fields, api
import json
import openpyxl
import base64
import re
import logging
from io import BytesIO
from openpyxl.utils import get_column_letter, column_index_from_string

_logger = logging.getLogger(__name__)

class ProductCategory(models.Model):
    _inherit = "product.category"
    
    template_file = fields.Binary(string="Upload Calculation Template")
    template_filename = fields.Char(string="Template Filename")
    
    # Google Sheets Integration
    google_sheet_url = fields.Char(string="Google Sheet System URL", help="Paste the full URL of the Google Sheet here.")
    
    
    spreadsheet_data = fields.Text(
        string="Spreadsheet Data",
        compute='_compute_spreadsheet_data',
        store=True
    )

    @api.onchange('google_sheet_url')
    def _onchange_google_sheet_url(self):
        """
        Reset the template file if the URL is cleared, allowing manual upload.
        """
        if not self.google_sheet_url:
            self.template_file = False
            self.template_filename = False

    
    @api.depends('template_file')
    def _compute_spreadsheet_data(self):
        for category in self:
            print("\n=========== DEBUG: _compute_spreadsheet_data ===========")
            print("CATEGORY:", category.name)
            print("template_file present?:", bool(category.template_file))
            
            if category.template_file:
                excel_data = category._convert_excel_to_spreadsheet(category.template_file)
                
                if excel_data:
                    category.spreadsheet_data = json.dumps(excel_data)
                    # small checksum for debugging
                    total_cells = sum(len(s.get('cells', {})) for s in excel_data.get('sheets', []))
                    total_merges = sum(len(s.get('merges', [])) for s in excel_data.get('sheets', []))
                    print("DEBUG: spreadsheet_data saved successfully! sheets:", len(excel_data.get('sheets', [])),
                          "cells:", total_cells, "merges:", total_merges)
                else:
                    print("DEBUG: excel_data conversion failed!")
            else:
                print("DEBUG: template_file is empty!")
                category.spreadsheet_data = False
            
            print("========================================================\n")
    
    def action_sync_google_sheet(self):
        """
        Fetches the XLSX directly from Google Drive URL and saves it to template_file.
        This triggers the existing _compute_spreadsheet_data logic.
        """
        self.ensure_one()
        if not self.google_sheet_url:
            return
            
        try:
            import requests
            from odoo.exceptions import UserError
            
            # 1. Parse URL to get Sheet ID
            # Formats: 
            # https://docs.google.com/spreadsheets/d/SPREADSHEET_ID/edit...
            # https://docs.google.com/spreadsheets/d/e/PUBLISHED_ID/pubhtml...
            
            sheet_id = False
            # Regex to capture ID between /d/ and /
            import re
            match = re.search(r'/d/([a-zA-Z0-9-_]+)', self.google_sheet_url)
            if match:
                sheet_id = match.group(1)
            
            if not sheet_id:
                raise UserError("Could not extract Spreadsheet ID. Please verify the URL format (should contain /d/spreadsheet_id/).")
                
            # 2. Construct Export URL
            export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
            
            _logger.info(f"Attempting to download Google Sheet from: {export_url}")
            
            # 3. Fetch Data
            response = requests.get(export_url, timeout=30)
            
            if response.status_code == 200:
                # 4. Save to binary field (this triggers the compute method)
                self.template_file = base64.b64encode(response.content)
                self.template_filename = "GoogleSheet_Synced.xlsx"
                
                # Retrieve the sheet name if possible
                if 'Content-Disposition' in response.headers:
                    fname = re.findall('filename="(.+)"', response.headers['Content-Disposition'])
                    if fname:
                        self.template_filename = fname[0]
                
                self._compute_spreadsheet_data()
                
                return {
                    'type': 'ir.actions.client',
                    'tag': 'display_notification',
                    'params': {
                        'title': 'Success',
                        'message': 'Google Sheet synced successfully!',
                        'type': 'success',
                        'sticky': False,
                        'next': {'type': 'ir.actions.act_window_close'},
                    }
                }
            elif response.status_code == 401:
                raise UserError("Permission Denied (401).\n\nReason: The Google Sheet is Private.\n\nSolution:\n1. Open Google Sheet.\n2. Click 'Share' (top right).\n3. Under 'General access', select 'Anyone with the link'.\n4. Try 'Sync Now' again.")
            else:
                raise UserError(f"Failed to download. Status Code: {response.status_code}\nURL: {export_url}")

                
        except Exception as e:
            _logger.error(f"Google Sync Failed: {e}")
            raise e

    

    def _parse_merge_range(self, range_str):
        """
        Convert 'A1:B3' -> dict {'top':0,'left':0,'bottom':2,'right':1}
        (0-based indices)
        """
        try:
            parts = range_str.split(':')
            if len(parts) == 1:
                # single cell treated as no-merge
                col = ''.join([c for c in parts[0] if c.isalpha()])
                row = ''.join([c for c in parts[0] if c.isdigit()])
                left = column_index_from_string(col) - 1
                top = int(row) - 1
                return {'top': top, 'left': left, 'bottom': top, 'right': left}
            start, end = parts
            col1 = ''.join([c for c in start if c.isalpha()])
            row1 = ''.join([c for c in start if c.isdigit()])
            col2 = ''.join([c for c in end if c.isalpha()])
            row2 = ''.join([c for c in end if c.isdigit()])
            left = column_index_from_string(col1) - 1
            top = int(row1) - 1
            right = column_index_from_string(col2) - 1
            bottom = int(row2) - 1
            return {'top': top, 'left': left, 'bottom': bottom, 'right': right}
        except Exception as e:
            # fallback: return None so caller can ignore
            print("DEBUG: parse_merge_range failed for", range_str, "error:", e)
            return None

    def _clean_sheet_reference(self, ref):
        """
        Convert 'resin $D$2:$D$22' to 'Resin!D2:D22'
        or 'Sheet1 $A$1:$A$10' to 'Sheet1!A1:A10'
        Handles both single references and comparison expressions
        """
        if not ref:
            return ref
            
        # Remove dollar signs first
        ref = ref.replace('$', '')
        
        # Check if there's a space (indicating sheet name + range)
        parts = ref.strip().split()
        if len(parts) >= 2:
            # parts[0] is sheet name (like 'resin'), parts[1] onwards is range/expression
            sheet_name = parts[0].strip()
            remaining = ' '.join(parts[1:]).strip()
            
            # Capitalize first letter of sheet name
            sheet_name = sheet_name.capitalize()
            
            # Check if remaining part has comparison operators
            if any(op in remaining for op in ['=', '<', '>', '!=']):
                # Split by comparison operator but keep the operator
                for op in ['!=', '<=', '>=', '=', '<', '>']:
                    if op in remaining:
                        range_part, value_part = remaining.split(op, 1)
                        return f"{sheet_name}!{range_part.strip()} {op} {value_part.strip()}"
            else:
                # Just a range reference
                return f"{sheet_name}!{remaining}"
        else:
            # No space, just return cleaned reference
            return ref.strip()

    def _convert_excel_formula_to_odoo(self, formula, sheet_name, cell_ref):
        """
        Convert Excel formulas with __xludf.dummyfunction wrapper to Odoo format.
        Handles:
        - __xludf.dummyfunction() wrapper removal
        - Triple-quoted strings (\"\"\"VALUE\"\"\")
        - IFERROR, UNIQUE, FILTER functions (Odoo supports these natively)
        - Array MATCH formulas with multiplication to FILTER format
        
        Returns: Converted formula string or None if conversion fails
        """
        if not formula:
            return None
        
        original = formula
        # _logger.info(f"🔧 [FORMULA CONVERT] Sheet: {sheet_name}, Cell: {cell_ref}")
        # _logger.info(f"   📥 Original: {original[:150]}")
        
        try:
            # Loop until all __xludf.dummyfunction wrappers are removed
            # CASE INSENSITIVE CHECK
            while "__xludf" in formula.lower() or "dummyfunction" in formula.lower():
                formula_lower = formula.lower()
                
                # Try to find start of function
                pattern_start = "__xludf.dummyfunction"
                start_idx = formula_lower.find(pattern_start)
                
                # If not found with prefix, check just dummyfunction (sometimes prefix is missing or different)
                if start_idx == -1:
                    pattern_start = "dummyfunction"
                    start_idx = formula_lower.find(pattern_start)
                    
                if start_idx == -1:
                    break
                
                # Find the opening parenthesis after the function name
                open_paren_idx = formula.find("(", start_idx)
                if open_paren_idx == -1:
                    break
                    
                # The content starts after ("
                # We expect the structure: FUNCTION(" CONTENT ")
                # So we look for the first quote after the opening parenthesis
                first_quote_idx = formula.find('"', open_paren_idx)
                
                if first_quote_idx == -1:
                    break
                    
                content_start_idx = first_quote_idx + 1
                current_idx = content_start_idx
                found_end = False
                
                # Scan forward to find the closing ") pair
                while current_idx < len(formula):
                    char = formula[current_idx]
                    
                    if char == '"':
                        # Check next char to see if it's an escaped quote ""
                        if current_idx + 1 < len(formula) and formula[current_idx + 1] == '"':
                            current_idx += 2
                            continue
                        # Check if it is the closing quote followed by closing parenthesis
                        elif current_idx + 1 < len(formula) and formula[current_idx + 1] == ')':
                            # Found the end!
                            content_end_idx = current_idx
                            end_wrapper_idx = current_idx + 2 # Skip ")
                            
                            # Extract content
                            raw_content = formula[content_start_idx:content_end_idx]
                            
                            # Unescape Excel double quotes: "" -> "
                            unescaped_content = raw_content.replace('""', '"')
                            
                            # Replace the whole wrapper with the inner content
                            # We replace from start_idx (function name) up to end_wrapper_idx
                            formula = formula[:start_idx] + unescaped_content + formula[end_wrapper_idx:]
                            
                            _logger.info(f"   🔄 Unwrapped (Robust): {unescaped_content[:80]}")
                            found_end = True
                            break
                        else:
                            # It's a quote not followed by quote or paren? 
                            # Could be end of string arg inside formula, but we typically look for ")
                            # Let's assume it's part of the content if we didn't match the closing pattern
                            current_idx += 1
                    else:
                        current_idx += 1
                
                if not found_end:
                    _logger.warning(f"   ⚠️ Could not find closing pattern for __xludf in {formula}")
                    # Break to avoid infinite loop
                    break

            # Step 2: Handle triple-quoted strings ("""VALUE""" -> VALUE) if any remain (legacy regex artifact, might not be needed but keeping for safety)
            if '"""' in formula:
                formula = re.sub(r'"""([^"]*?)"""', r'\1', formula)
                _logger.info(f"   🔄 Fixed triple quotes")
            
            # Step 3: Ensure formula starts with =
            if not formula.startswith('='):
                formula = '=' + formula
            
            # Step 4: Ensure no double equals
            if formula.startswith('=='):
                formula = formula[1:]
            
            # Step 5: Clean up any extra quotes that might break syntax
            # Be careful not to break string literals inside formulas
            # checking for empty quotes '' might be dangerous if it's an empty string ""
            # formula = formula.replace("''", "'") 
            
            if original != formula:
                _logger.info(f"   ✅ Converted: {formula}")
            
            # -------------------------------------------------------------------------
            # SPECIAL HANDLING for array formulas like INDEX(.... MATCH(1, (cond1)*(cond2), 0))
            # which Odoo's engine often fails to evaluate (it doesn't do array mult in MATCH).
            # We convert this to: INDEX(FILTER(range, (cond1)*(cond2)), 1)
            # -------------------------------------------------------------------------
            if 'MATCH' in formula and '*' in formula:
                _logger.info(f"   🧐 Checking for array MATCH pattern in: {formula}")
                
                # Regex to capture: INDEX( (return_range), MATCH(1, (criteria_expr), 0) )
                # Note: This is an approximation. Nested parens can be tricky for regex.
                # structure: INDEX( \s* ([^,]+) \s*, \s* MATCH \s* \( \s* 1 \s*, \s* (.+?) \s*, \s* 0 \s* \) \s* \)
                # Added \s* after MATCH to handle spaces before (
                
                # We'll try a flexible pattern. 
                # Group 1: Return Range
                # Group 2: Criteria Expression (the (A=1)*(B=2) part)
                pattern = r"INDEX\(\s*([^,]+)\s*,\s*MATCH\s*\(\s*1\s*,\s*(.+?)\s*,\s*0\s*\)\s*\)"
                
                match = re.search(pattern, formula, re.IGNORECASE)
                if match:
                    return_range = self._clean_sheet_reference(match.group(1).strip())
                    criteria_expr = match.group(2).strip()
                    
                    # Clean sheet references in criteria expression
                    # The criteria looks like: (resin $A$2:$A$22 = L3) * (resin $B$2:$B$22 = M3)
                    # We need to convert each part separately
                    
                    # Split by * operator while preserving parentheses structure
                    # Use regex to split on * that's outside parentheses
                    parts = re.split(r'\)\s*\*\s*\(', criteria_expr)
                    
                    cleaned_parts = []
                    for i, part in enumerate(parts):
                        # Add back parentheses that were removed by split
                        if i == 0:
                            # First part - remove leading ( if exists
                            part = part.lstrip('(')
                        elif i == len(parts) - 1:
                            # Last part - remove trailing ) if exists
                            part = part.rstrip(')')
                        
                        # Clean the sheet reference in this condition
                        cleaned_part = self._clean_sheet_reference(part.strip())
                        cleaned_parts.append(cleaned_part)
                    
                    # Join with FILTER's comma separator and wrap each in parentheses
                    new_criteria = '), ('.join(cleaned_parts)
                    
                    # Construct valid Odoo spreadsheet formula: INDEX(FILTER(return_range, cond1, cond2...), 1)
                    new_formula = f"INDEX(FILTER({return_range}, ({new_criteria})), 1)"
                    
                    # If wrapped in IFERROR, preserve it.
                    # The original might be =IFERROR(INDEX(...), "")
                    # We just replaced the INDEX(...) part.
                    # We can use simple string replacement of the matched span
                    
                    start_idx, end_idx = match.span()
                    final_formula = formula[:start_idx] + new_formula + formula[end_idx:]
                    
                    _logger.info(f"   🚀 CONVERTED ARRAY MATCH: {final_formula}")
                    return final_formula

            return formula
            
        except Exception as e:
            _logger.error(f"   ❌ Conversion failed: {e}", exc_info=True)
            return None


    def _convert_excel_to_spreadsheet(self, file_data):
        """
        Convert uploaded XLSX (binary base64) into Odoo Spreadsheet JSON structure.
        Returns dict or None on failure.
        """
        try:
            print("\n******** DEBUG: Converting XLSX → Odoo Spreadsheet (openpyxl) ********")

            # decode and load workbook
            file_content = base64.b64decode(file_data)
            wb = openpyxl.load_workbook(BytesIO(file_content), data_only=False)

            spreadsheet = {
                "version": 16,
                "sheets": [],
                "revisionId": 1,
                "settings": {},
                "lists": {},
                "formats": {},
                "styles": {},
                "borders": {},
            }

            for sheet in wb.worksheets:
                # determine dimensions (fallbacks)
                max_row = sheet.max_row or 1
                max_col = sheet.max_column or 1

                # build sheet_json
                sheet_json = {
                    # unique id to avoid collision with sheet_<line.id>: prefix template_
                    "id": ("template_" + (sheet.title or "Sheet")).replace(" ", "_")[:60],
                    "name": (sheet.title or "Sheet")[:31],
                    # colNumber/rowNumber : use counts (Odoo expects integer)
                    "colNumber": int(max_col),
                    "rowNumber": max(int(max_row), 1000) if "profile master" in (sheet.title or "").lower() else int(max_row),
                    "cells": {},
                    "merges": [],
                    "rows": {},  # numeric-string keys: "0","1"
                    "cols": {},  # numeric-string keys: "0","1"
                }

                # -------------------------
                # merges: convert to numeric boxes
                # -------------------------
                try:
                    for merged in getattr(sheet, "merged_cells").ranges:
                        rng = str(merged)  # like 'A1:B3'
                        parsed = self._parse_merge_range(rng)
                        if parsed:
                            sheet_json["merges"].append(parsed)
                except Exception as e:
                    # if no merges or failure, ignore but log
                    print("DEBUG: reading merges failed for sheet", sheet.title, "error:", e)

                # -------------------------
                # column widths -> Odoo expects numeric index keys as strings
                # openpyxl.column_dimensions keys are letters like 'A'
                # convert: letter -> index-1 -> string key
                # -------------------------
                try:
                    for col_letter, col_dim in sheet.column_dimensions.items():
                        width = getattr(col_dim, "width", None)
                        if width is not None:
                            try:
                                idx = column_index_from_string(col_letter) - 1
                                sheet_json["cols"][str(idx)] = {"width": float(width)}
                            except Exception:
                                continue
                except Exception as e:
                    print("DEBUG: reading column_dimensions failed for sheet", sheet.title, "error:", e)

                # -------------------------
                # row heights -> numeric-string keys
                # -------------------------
                try:
                    for r_idx, row_dim in sheet.row_dimensions.items():
                        height = getattr(row_dim, "height", None)
                        if height is not None:
                            try:
                                sheet_json["rows"][str(int(r_idx) - 1)] = {"size": float(height)}
                            except Exception:
                                continue
                except Exception as e:
                    print("DEBUG: reading row_dimensions failed for sheet", sheet.title, "error:", e)

                # -------------------------
                # cells: iterate full rectangle so positions align
                # -------------------------
                _logger.info(f"\n📊 Processing sheet: {sheet.title} ({max_row} rows x {max_col} cols)")
                formula_count = 0
                xludf_count = 0
                converted_count = 0
                skipped_count = 0
                
                for r in range(1, max_row + 1):
                    for c in range(1, max_col + 1):
                        cell = sheet.cell(row=r, column=c)
                        if cell is None:
                            continue
                        if cell.value is None:
                            # skip fully empty cells
                            continue

                        # key as A1 etc. Odoo expects A1-style keys inside cells dict
                        col_letter = get_column_letter(c)
                        key = f"{col_letter}{r}"
                        
                        # Handle ArrayFormula objects (can appear in any cell type)
                        from openpyxl.worksheet.formula import ArrayFormula
                        if isinstance(cell.value, ArrayFormula):
                            # Extract formula text from ArrayFormula object
                            raw = cell.value.text if hasattr(cell.value, 'text') else ""
                            if raw:
                                formula_count += 1
                                content = raw if raw.startswith('=') else '=' + raw
                            else:
                                # Empty array formula - skip
                                continue
                        # decide content: formula vs value
                        elif cell.data_type == 'f':
                            # Regular formula
                            raw = str(cell.value) if cell.value is not None else ""
                            formula_count += 1
                            
                            # 🔥 Process formulas with __xludf - convert Excel dynamic arrays to Odoo format
                            if '__xludf' in raw or '__xlud' in raw:
                                xludf_count += 1
                                _logger.info(f"🔍 Found __xludf formula in {sheet.title} at {key}")
                                
                                # Convert to Odoo format
                                converted = self._convert_excel_formula_to_odoo(raw, sheet.title, key)
                                if converted:
                                    content = converted
                                    converted_count += 1
                                    _logger.info(f"✅ Successfully converted formula at {key}")
                                else:
                                    skipped_count += 1
                                    _logger.warning(f"❌ Failed to convert formula at {key}, skipping")
                                    continue
                            # 🔥 Check for Array MATCH formulas even without __xludf (e.g. manually entered or pre-existing)
                            elif 'MATCH' in raw and '*' in raw:
                                # Start with assuming it might need conversion
                                _logger.info(f"🔍 Found potential Array MATCH formula in {sheet.title} at {key}: {raw[:100]}...")
                                converted = self._convert_excel_formula_to_odoo(raw, sheet.title, key)
                                if converted and converted != raw:
                                    content = converted
                                    converted_count += 1
                                    _logger.info(f"✅ Successfully converted Array MATCH formula at {key}")
                                else:
                                    # If no change, just use raw
                                    content = raw if raw.startswith('=') else '=' + raw
                            else:
                                content = raw if raw.startswith('=') else '=' + raw
                        else:
                            # preserve native python types for numbers/bool
                            v = cell.value
                            # Skip error values
                            if isinstance(v, str) and v.startswith('#'):
                                continue
                            # openpyxl may return datetime objects for dates — keep them as isoformat strings
                            try:
                                import datetime
                                if isinstance(v, (datetime.date, datetime.datetime)):
                                    content = v.isoformat()
                                else:
                                    content = v
                            except Exception:
                                content = str(v)
                        sheet_json["cells"][key] = {
                            "content": str(content) if content is not None else "",
                        }

                        # 🔥 FIX for #SPILL! errors: 
                        # Skip cells that contain valid formula artifacts like "COMPUTED_VALUE" or "#SPILL!"
                        # Odoo will just overwrite these empty cells when the array formula expands.
                        # 🔥 FIX for #SPILL! errors: 
                        # Skip cells that contain valid formula artifacts like "COMPUTED_VALUE" or "#SPILL!"
                        # Odoo will just overwrite these empty cells when the array formula expands.
                        content_str = str(content).strip()
                        c_upper = content_str.upper()
                        
                        # Check for various artifact patterns
                        is_artifact = (
                            c_upper == "COMPUTED_VALUE" or
                            "COMPUTED_VALUE" in c_upper or  # Catch _xlfn...COMPUTED_VALUE
                            c_upper.startswith("#SPILL") or 
                            c_upper.startswith("#CALC") or
                            c_upper == "#N/A" # Sometimes array formulas start with N/A before calculation
                        )
                        
                        if is_artifact:
                            _logger.info(f"   🧹 Skipping artifact in {key}: {content}")
                            del sheet_json["cells"][key]
                            continue

                        # 🔥 SPECIAL FIX for 'Helper' Sheet Spill Zones
                        # The user has dynamic formulas in Columns A, C, E (Indices 1, 3, 5) starting at Row 2.
                        # Google Sheets exports the calculated results as static values in rows 3+, which block the spill.
                        # We must clear these static values to allow the Odoo formulas to expand.
                        # We assume Row 2 contains the formula, so we skip clearing Row 2.
                        if "helper" in sheet.title.lower() and c in [1, 3, 5] and r > 2:
                             # If it's not a formula (cell.data_type != 'f'), it's likely a blocker artifact.
                             # Note: We already processed cell.data_type='f' above and would have set content accordingly.
                             # But here we are looking at the *source* cell type.
                             # We can check if we detected a formula in this cell.
                             # If key is in sheet_json['cells'] and content does not start with '=', it's a value.
                             
                             current_content = sheet_json["cells"][key]["content"]
                             if not str(current_content).startswith("="):
                                 _logger.info(f"   🧹 Skipping Helper Sheet Spill Blockage in {key}: {current_content}")
                                 del sheet_json["cells"][key]
                                 continue


                
                # Log summary for this sheet
                _logger.info(f"📈 Sheet '{sheet.title}' processing summary:")
                _logger.info(f"   Total formulas: {formula_count}")
                _logger.info(f"   __xludf formulas found: {xludf_count}")
                _logger.info(f"   Successfully converted: {converted_count}")
                _logger.info(f"   Skipped/failed: {skipped_count}")

                # -------------------------
                # Data Validations (Dropdowns)
                # -------------------------
                try:
                    if hasattr(sheet, 'data_validations') and sheet.data_validations:
                        sheet_json['validations'] = []
                        for dv in sheet.data_validations.dataValidation:
                            # We only care about LIST type for dropdowns usually
                            if dv.type == 'list':
                                # dv.sqref is a generic 'A1:A10 B1:B10' string or MultiCellRange
                                # We'll store it as a list of ranges
                                ranges = str(dv.sqref).split()
                                sheet_json['validations'].append({
                                    'type': 'list',
                                    'formula1': dv.formula1,
                                    'ranges': ranges,
                                    'showErrorMessage': dv.showErrorMessage,
                                    'showInputMessage': dv.showInputMessage,
                                })
                                print(f"DEBUG: Found validation in {sheet.title} | Type: {dv.type} | Formula: {dv.formula1} | Ranges: {ranges}")
                            else:
                                print(f"DEBUG: Ignored validation type: {dv.type} in {sheet.title}")
                except Exception as e:
                    print("DEBUG: reading data_validations failed for sheet", sheet.title, "error:", e)

                # append sheet
                spreadsheet["sheets"].append(sheet_json)
                print("DEBUG: parsed sheet:", sheet.title, "cells:", len(sheet_json["cells"]),
                      "merges:", len(sheet_json["merges"]),
                      "cols_meta:", len(sheet_json["cols"]),
                      "rows_meta:", len(sheet_json["rows"]))

            print("DEBUG: XLSX parsed. Total sheets:", len(spreadsheet["sheets"]))
            return spreadsheet

        except Exception as e:
            print("DEBUG: Excel conversion failed:", e)
            return None